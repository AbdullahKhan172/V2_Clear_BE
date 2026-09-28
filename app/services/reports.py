"""
Report generation - the downloadable deliverables.
==================================================
Four formats, all built from the same ProjectDataModel the dashboard uses:

    html    self-contained interactive dashboard (charts, 3D, methodology)
    excel   full BOQ with the emission factors applied
    word    A3 landscape summary report, ready to issue
    boq     extracted quantities, showing what was USED vs EXCLUDED

The first three delegate to the existing, verified generators in reports/. Only
`boq` is moved here, from web_app.py:1230-1355, because it was built inline in
the route and exists nowhere else.

Why the HTML report is kept even though a React dashboard is coming: it is the
artefact that survives being emailed. It works offline, needs no server, and
still opens in five years. The React dashboard is the interactive view; this is
the deliverable.
"""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd

from app import blobs

# kind -> (extension, mime type, human label)
REPORT_KINDS = {
    'html': ('html', 'text/html', 'Interactive dashboard'),
    'excel': ('xlsx',
              'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
              'Excel workbook'),
    'word': ('docx',
             'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
             'Word report'),
    'boq': ('xlsx',
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'Quantities export'),
}


def safe_name(value: str, fallback: str = 'report', max_len: int = 60) -> str:
    """Filesystem- and header-safe name. web_app.py:126, moved.

    Used for the download filename, so a project name can never smuggle a path
    separator or CRLF header injection into Content-Disposition.
    """
    s = re.sub(r'[^\w\-]', '_', str(value or '')).strip('_')
    return (s[:max_len] or fallback)


def generate(run_id: str, kind: str, result, *, project_name: str,
             full_elements_df: pd.DataFrame | None = None,
             filtered_elements_df: pd.DataFrame | None = None) -> dict:
    """Build one report and store it as a blob. Returns its manifest entry.

    Args:
        result:                RunResult from execute_run
        full_elements_df:      every extracted element, BEFORE Step-2 filtering
        filtered_elements_df:  what survived - `boq` needs both, to show what was
                               excluded as well as what was counted
    """
    if kind not in REPORT_KINDS:
        raise ValueError(f'Unknown report kind {kind!r}; '
                         f'choose from {sorted(REPORT_KINDS)}.')
    ext, mime, label = REPORT_KINDS[kind]

    key = f'{run_id}/reports/{kind}.{ext}'
    target = blobs.BLOB_ROOT / key
    target.parent.mkdir(parents=True, exist_ok=True)

    if kind == 'html':
        from reports.html_report import generate_html_report
        generate_html_report(result.data, str(target), 'advanced')
    elif kind == 'excel':
        from reports.excel_report import generate_excel_report
        generate_excel_report(result.data, str(target))
    elif kind == 'word':
        from reports.word_report import generate_word_report
        generate_word_report(result.data, str(target))
    else:
        _build_quantities_export(target, full_elements_df, filtered_elements_df)

    from datetime import datetime, timezone
    return {
        'kind': kind,
        'key': key,
        'label': label,
        'ext': ext,
        'mime': mime,
        'filename': f'{safe_name(project_name)}_{kind}.{ext}',
        'size': target.stat().st_size,
        'generated_at': datetime.now(timezone.utc).isoformat(),
    }


def _build_quantities_export(target: Path, full_df, filt_df) -> None:
    """The "Quantities export" workbook. web_app.py:1230-1355, moved verbatim.

    Shows every extracted element with its total vs used quantity, so a reader
    can see what was left out of the assessment as well as what went in - which
    is the first question anyone reviewing a carbon figure asks.
    """
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    if full_df is None:
        full_df = pd.DataFrame()
    if filt_df is None:
        filt_df = pd.DataFrame()

    cat_col = next((c for c in ('Category', 'category') if c in full_df.columns), None)
    name_col = next((c for c in ('Description', 'name') if c in full_df.columns), None)
    vol_col = next((c for c in ('Volume(m3)', 'volume_m3') if c in full_df.columns), None)
    cnt_col = next((c for c in ('Count', 'count') if c in full_df.columns), None)

    def _safe(row, col, default=''):
        if col and col in row.index:
            v = row[col]
            return default if (v != v) else v          # NaN check
        return default

    def _num(row, col):
        try:
            return float(_safe(row, col, 0))
        except (TypeError, ValueError):
            return 0.0

    def _mass(row):
        # IFC extraction stores mass under 'weight_kg'; BOQ/CSV under 'Mass(kg)'.
        # Falling back means IFC steel sections report their real modelled mass
        # instead of 0 kg.
        try:
            return float(row.get('Mass(kg)', 0) or row.get('weight_kg', 0) or 0)
        except (TypeError, ValueError):
            return 0.0

    groups: dict = {}
    for _, row in full_df.iterrows():
        key = (str(_safe(row, cat_col)), str(_safe(row, name_col)))
        g = groups.setdefault(key, {'cat': key[0], 'name': key[1], 'total_vol': 0.0,
                                    'total_mass': 0.0, 'total_cnt': 0})
        g['total_vol'] += _num(row, vol_col)
        g['total_mass'] += _mass(row)
        try:
            g['total_cnt'] += int(_num(row, cnt_col) or 1)
        except (TypeError, ValueError):
            pass

    used: dict = {}
    for _, row in filt_df.iterrows():
        key = (str(_safe(row, cat_col)), str(_safe(row, name_col)))
        u = used.setdefault(key, {'vol': 0.0, 'mass': 0.0, 'cnt': 0})
        u['vol'] += _num(row, vol_col)
        u['mass'] += _mass(row)
        try:
            u['cnt'] += int(_num(row, cnt_col) or 1)
        except (TypeError, ValueError):
            pass

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'BOQ Quantities'

    hdr_fill = PatternFill('solid', fgColor='1E3A5F')
    hdr_font = Font(color='FFFFFF', bold=True, size=10)
    used_fill = PatternFill('solid', fgColor='F0FDF4')
    excl_fill = PatternFill('solid', fgColor='FEF2F2')
    partial_fill = PatternFill('solid', fgColor='FFFBEB')

    headers = ['Category', 'Element Name', 'Total Volume (m³)', 'Used Volume (m³)',
               'Excluded Volume (m³)', 'Total Mass (kg)', 'Used Mass (kg)',
               'Total Count', 'Used Count', 'Status']
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal='center')

    for key, g in sorted(groups.items()):
        u = used.get(key, {'vol': 0.0, 'mass': 0.0, 'cnt': 0})
        total_vol = round(g['total_vol'], 3)
        used_vol = round(u['vol'], 3)
        excl_vol = round(total_vol - used_vol, 3)

        # A row counts as "Used" if it contributed volume OR mass - steel and PT
        # have no volume but are quantified, and emitted, by mass.
        if used_vol + round(u['mass'], 1) <= 0:
            status, fill = 'Excluded', excl_fill
        elif excl_vol > 0.001:
            status, fill = 'Partial (level/manual filter)', partial_fill
        else:
            status, fill = 'Used', used_fill

        ws.append([g['cat'], g['name'], total_vol, used_vol, excl_vol,
                   round(g['total_mass'], 1), round(u['mass'], 1),
                   g['total_cnt'], u['cnt'], status])
        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=ws.max_row, column=col_idx).fill = fill

    for col, width in zip('ABCDEFGHIJ', [20, 45, 16, 14, 16, 16, 14, 12, 10, 30]):
        ws.column_dimensions[col].width = width

    wb.save(str(target))
